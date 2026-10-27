from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

ROOT = Path(r"C:\Users\Lenovo LEGION 5\OneDrive\Desktop\VISTA")
OUT = ROOT / "outputs" / "Cam_nang_trien_khai_LLM_tu_host_cho_VISTA.docx"
OUT.parent.mkdir(exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.82)
sec.right_margin = Inches(0.82)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
styles['Normal'].font.size = Pt(10.4)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.13
for name, size in [('Title', 28), ('Subtitle', 13), ('Heading 1', 18), ('Heading 2', 14), ('Heading 3', 11.5)]:
    s = styles[name]
    s.font.name = 'Aptos Display' if name != 'Normal' else 'Aptos'
    s._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos Display')
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.bold = name != 'Subtitle'
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
    s.paragraph_format.space_after = Pt(6)

if 'Code' not in styles:
    cs = styles.add_style('Code', WD_STYLE_TYPE.PARAGRAPH)
else:
    cs = styles['Code']
cs.font.name = 'Consolas'
cs._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
cs.font.size = Pt(8.3)
cs.paragraph_format.left_indent = Inches(0.22)
cs.paragraph_format.right_indent = Inches(0.12)
cs.paragraph_format.space_before = Pt(4)
cs.paragraph_format.space_after = Pt(6)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margin(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'),'dxa')

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'17365D'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
        set_cell_margin(c)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cells[i], 'F2F6FA' if ri%2 else 'FFFFFF'); set_cell_margin(cells[i])
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs: r.font.size=Pt(8.7)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def p(text='', bold_lead=None):
    para=doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        para.add_run(bold_lead).bold=True; para.add_run(text[len(bold_lead):])
    else: para.add_run(text)
    return para

def bullet(text, level=0):
    para=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2')
    para.add_run(text); return para

def number(text, level=0):
    para=doc.add_paragraph(style='List Number' if level==0 else 'List Number 2')
    para.add_run(text); return para

def code(text):
    para=doc.add_paragraph(style='Code'); para.add_run(text); return para

def h1(text): doc.add_heading(text, level=1)
def h2(text): doc.add_heading(text, level=2)
def h3(text): doc.add_heading(text, level=3)
def page(): doc.add_page_break()

# Cover
title=doc.add_paragraph(style='Title'); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
title.add_run('Cẩm nang triển khai LLM tự host cho hệ thống VISTA')
sub=doc.add_paragraph(style='Subtitle'); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
sub.add_run('Kiến thức nền tảng  Kiến trúc  RAG  Bảo mật  Vận hành  Lộ trình triển khai')
doc.add_paragraph()
intro=doc.add_paragraph(); intro.alignment=WD_ALIGN_PARAGRAPH.CENTER
intro.add_run('Phiên bản 1.0  |  Tháng 10 năm 2026').bold=True
doc.add_paragraph()
p('Tài liệu này hướng dẫn đội VISTA chuyển từ việc phụ thuộc trực tiếp vào Gemini API sang một nền tảng AI do dự án tự kiểm soát. Kết luận chính là: không huấn luyện một mô hình ngôn ngữ lớn từ đầu; hãy tự host mô hình open weight, bổ sung dữ liệu VISTA bằng RAG, giữ toàn bộ luật nghiệp vụ trong Spring Boot và xây lớp AI Provider để có thể đổi mô hình mà không viết lại sản phẩm.')
p('Đối tượng đọc gồm thành viên backend, mobile, web, DevOps, kiểm thử và người quản lý kỹ thuật. Sau khi đọc, nhóm có thể chọn cấu hình phần cứng, dựng môi trường thử nghiệm, tích hợp model server, xây kho tri thức, bảo vệ dữ liệu, đánh giá chất lượng và chuyển đổi dần khỏi Gemini mà không làm gián đoạn dịch vụ.')
page()

h1('Mục lục nội dung')
for item in [
'1  Điều cần hiểu trước khi bắt đầu','2  LLM hoạt động như thế nào','3  Các cách xây AI và lựa chọn phù hợp','4  Kiến trúc mục tiêu cho VISTA','5  Lựa chọn mô hình và runtime','6  Lựa chọn phần cứng','7  RAG và kho tri thức du lịch','8  Embedding và tìm kiếm hybrid','9  Tool calling và luật nghiệp vụ','10  Structured output và kiểm chứng','11  Thiết kế AI service','12  Tích hợp Spring Boot','13  Hàng đợi và streaming','14  Bảo mật và quyền riêng tư','15  Quan sát và kiểm soát chi phí','16  Bộ đánh giá chất lượng','17  Fine tuning và LoRA','18  Quy trình triển khai từng giai đoạn','19  Runbook vận hành','20  Checklist nghiệm thu']:
    p(item)

h1('1  Điều cần hiểu trước khi bắt đầu')
h2('1.1  Mục tiêu thực tế')
p('Mục tiêu của VISTA không phải là tạo một mô hình nền tảng mới cạnh tranh với các hãng AI. Mục tiêu là sở hữu chuỗi xử lý AI: chọn model, tự vận hành inference, kiểm soát dữ liệu, kiểm soát quota, gắn kiến thức du lịch của VISTA và thay model khi cần.')
p('Một hệ thống AI production gồm nhiều thành phần hơn model. Model chỉ sinh token. Chất lượng sản phẩm còn phụ thuộc vào dữ liệu truy xuất, prompt, validation, công cụ nghiệp vụ, bảo mật, quan sát, giao diện và quy trình xử lý lỗi.')
h2('1.2  Bốn khái niệm thường bị nhầm')
table(['Khái niệm','Ý nghĩa','VISTA nên dùng khi nào'],[
('Pretraining','Huấn luyện hàng tỷ tham số từ kho dữ liệu cực lớn.','Không phù hợp cho dự án hiện tại.'),
('Fine tuning','Điều chỉnh hành vi model bằng tập ví dụ có nhãn.','Chỉ cân nhắc sau khi có dữ liệu đánh giá tốt.'),
('RAG','Tìm dữ liệu liên quan rồi đưa vào context trước khi sinh câu trả lời.','Là nền tảng chính để dùng dữ liệu du lịch VISTA.'),
('Prompt engineering','Thiết kế instruction, schema và ví dụ cho model.','Dùng ngay, nhưng phải có version và test.')], [1.2,2.7,2.7])
h2('1.3  Nguyên tắc nguồn sự thật')
bullet('LLM không phải nguồn sự thật cho giá, giờ mở cửa, số dư, quyền truy cập hoặc trạng thái thanh toán.')
bullet('PostgreSQL và các service nghiệp vụ là nguồn sự thật.')
bullet('LLM chỉ đề xuất; backend xác minh rồi mới lưu.')
bullet('Khi không có dữ liệu đáng tin cậy, hệ thống phải nói không biết thay vì bịa.')

h1('2  LLM hoạt động như thế nào')
h2('2.1  Token và context')
p('LLM nhận văn bản đã được chia thành token và dự đoán token tiếp theo. Context window là tổng số token model có thể xử lý trong một request, bao gồm system prompt, lịch sử, dữ liệu RAG, tool result và output. Context dài làm tăng bộ nhớ KV cache và độ trễ; vì vậy không nên đẩy toàn bộ database vào prompt.')
h2('2.2  Tham số và lượng tử hóa')
p('Số tham số phản ánh kích thước model, không đảm bảo trực tiếp chất lượng cho mọi nhiệm vụ. Quantization chuyển trọng số sang độ chính xác thấp hơn như 8 bit hoặc 4 bit để giảm VRAM và tăng khả năng chạy trên phần cứng nhỏ. Đổi lại, chất lượng hoặc tốc độ có thể suy giảm tùy model và runtime.')
table(['Dạng','Ưu điểm','Hạn chế'],[
('BF16 hoặc FP16','Chất lượng ổn định, hỗ trợ tốt trên GPU server.','Tốn VRAM.'),
('INT8','Giảm bộ nhớ, thường giữ chất lượng khá tốt.','Không phải runtime nào cũng tối ưu như nhau.'),
('INT4 AWQ hoặc GPTQ','Phù hợp GPU 16 đến 24 GB.','Cần benchmark chất lượng và kernel.'),
('GGUF Q4 hoặc Q5','Dễ chạy với llama.cpp trên CPU hoặc GPU cá nhân.','Thông lượng production thường thấp hơn cấu hình GPU chuyên dụng.')],[1.4,2.7,2.5])
h2('2.3  Sampling')
p('Temperature, top p và max tokens ảnh hưởng tính ngẫu nhiên. Với lập lịch trình có JSON schema, nên dùng temperature thấp để ổn định. Không dùng temperature cao để sửa vấn đề thiếu dữ liệu; điều đó chỉ làm câu trả lời đa dạng hơn, không đáng tin hơn.')

h1('3  Các cách xây AI và lựa chọn phù hợp')
table(['Phương án','Chi phí đầu tư','Kiểm soát','Khuyến nghị'],[
('API bên ngoài','Thấp lúc đầu','Thấp đến trung bình','Giữ làm fallback tạm thời.'),
('Self host open weight','Trung bình','Cao','Phương án mục tiêu cho VISTA.'),
('Fine tune model có sẵn','Trung bình đến cao','Cao','Thực hiện sau RAG và evaluation.'),
('Pretrain từ đầu','Rất cao','Rất cao','Không triển khai ở giai đoạn này.')],[1.5,1.4,1.2,2.5])
p('Quyết định đề xuất là triển khai self hosted inference với interface tương thích OpenAI, để backend có thể gọi model local, model cloud hoặc Gemini qua cùng một abstraction. Điều này tránh khóa nhà cung cấp và cho phép chuyển traffic từng phần.')

h1('4  Kiến trúc mục tiêu cho VISTA')
code('Web và Mobile\n      |\n      v\nSpring Boot API\n  auth | TripAccess | quota | audit | validation\n      |\n      v\nVISTA AI Orchestrator\n  prompt | RAG | tool calling | JSON schema | guardrails\n      |                         |\n      v                         v\nModel server                Embedding service\nQwen3 qua vLLM             BGE M3\n      |                         |\n      +------------+------------+\n                   v\n        PostgreSQL và pgvector')
h2('4.1  Phân chia trách nhiệm')
table(['Thành phần','Chịu trách nhiệm','Không được làm'],[
('Web và Mobile','Thu thập yêu cầu, hiển thị tiến trình, lỗi và citations.','Không chứa API key hoặc tự xác nhận dữ liệu AI.'),
('Spring Boot','Auth, quyền chuyến đi, quota, transaction, validation, lưu dữ liệu.','Không nhúng model runtime vào cùng JVM.'),
('AI Orchestrator','RAG, prompt, gọi model, schema, tool loop.','Không truy cập database bằng SQL tùy ý.'),
('Model server','Inference và batching.','Không expose trực tiếp ra Internet.'),
('PostgreSQL','Nguồn dữ liệu thật và vector store.','Không cho model quyền database trực tiếp.')],[1.35,2.8,2.45])
h2('4.2  Luồng tạo lịch trình')
for x in ['Client gửi yêu cầu tới Spring Boot.','Spring Boot xác thực user và quyền chuyến đi.','Backend chuẩn hóa ngân sách, số ngày và sở thích.','Orchestrator truy xuất destination PUBLISHED bằng RAG.','Model lập phương án từ danh sách được phép.','Backend tính lại chi phí và kiểm tra mọi destination ID.','Nếu hợp lệ, trả preview; người dùng xác nhận mới lưu.']:
    number(x)

h1('5  Lựa chọn mô hình và runtime')
h2('5.1  Model khởi đầu')
p('Khuyến nghị bắt đầu với Qwen3 8B vì kích thước vừa phải, hỗ trợ tiếng Việt, có giấy phép Apache 2.0 và được hỗ trợ bởi các runtime phổ biến. Nếu bộ đánh giá cho thấy 8B chưa đạt, thử Qwen3 14B trước khi nhảy lên model rất lớn.')
table(['Môi trường','Model','Runtime','Mục tiêu'],[
('Laptop CPU hoặc GPU nhỏ','Qwen3 4B hoặc 8B GGUF Q4','llama.cpp','Phát triển chức năng.'),
('GPU 16 đến 24 GB','Qwen3 8B quantized','vLLM','MVP production ít tải.'),
('GPU 24 đến 48 GB','Qwen3 14B quantized','vLLM','Tăng chất lượng.'),
('Nhiều GPU','Qwen3 32B hoặc model thắng evaluation','vLLM distributed','Scale sau khi có tải thật.')],[1.45,1.6,1.4,2.1])
h2('5.2  llama.cpp và vLLM')
p('llama.cpp phù hợp môi trường phát triển vì có thể chạy GGUF trên CPU, CUDA, Vulkan và kết hợp CPU GPU. vLLM phù hợp server GPU vì có continuous batching, quản lý KV cache và parallelism tốt hơn. Cả hai có API tương thích OpenAI, giúp giữ nguyên contract giữa backend và model server.')
h2('5.3  Quy trình chọn model')
bullet('Pin chính xác model revision và checksum.')
bullet('Đọc giấy phép model và điều kiện phân phối.')
bullet('Chạy cùng một bộ prompt tiếng Việt trên từng ứng viên.')
bullet('Đo schema pass rate, hallucination, budget violation, latency và VRAM.')
bullet('Chọn model nhỏ nhất đạt ngưỡng thay vì model lớn nhất có thể chạy.')

h1('6  Lựa chọn phần cứng')
h2('6.1  Ước lượng bộ nhớ')
p('Bộ nhớ inference không chỉ chứa weights. Nó còn chứa KV cache, activation, runtime overhead và batch đang chạy. Một model vừa VRAM ở batch 1 chưa chắc chịu được nhiều người dùng hoặc context dài.')
table(['Yếu tố','Làm tăng tài nguyên','Cách kiểm soát'],[
('Model size','Nhiều tham số hơn.','Chọn model nhỏ hơn hoặc quantization.'),
('Context length','KV cache tăng theo token và concurrency.','Giới hạn context, RAG gọn và truncate có chủ đích.'),
('Concurrency','Nhiều sequence đồng thời.','Queue, admission control và batching.'),
('Output length','GPU bị giữ lâu hơn.','Đặt max tokens theo use case.'),
('Embedding jobs','Dùng CPU hoặc GPU song song.','Tách service và chạy batch ngoài giờ cao điểm.')],[1.5,2.55,2.1])
h2('6.2  Không mua GPU trước benchmark')
p('Hãy đo request mỗi ngày, concurrency đỉnh, context trung bình, output trung bình, time to first token và tokens mỗi giây. Sau đó thuê GPU theo giờ để benchmark trước khi mua. GPU 24 GB thường là điểm bắt đầu hợp lý cho model 8B, nhưng đây không phải bảo đảm; kết quả phụ thuộc quantization, context và tải đồng thời.')

h1('7  RAG và kho tri thức du lịch')
h2('7.1  Vì sao VISTA cần RAG')
p('Giá vé, giờ mở cửa và trạng thái địa điểm thay đổi. Fine tuning không phải cơ chế cập nhật dữ liệu thời gian thực. RAG cho phép model nhận đúng dữ liệu mới nhất ở mỗi request và trả kèm nguồn.')
h2('7.2  Pipeline ingestion')
for x in ['Đọc bản ghi destination và tài liệu đã được duyệt.','Chuẩn hóa Unicode, đơn vị tiền, thời gian và địa danh.','Tách chunk theo ngữ nghĩa, không cắt giữa trường quan trọng.','Gắn metadata và content hash.','Tạo embedding theo batch.','Upsert chunk và embedding.','Đánh dấu lại khi source thay đổi hoặc bị unpublish.']:
    number(x)
h2('7.3  Schema dữ liệu')
code('ai_knowledge_documents\n  id, source_type, source_id, title, content_hash, language, status, updated_at\n\nai_knowledge_chunks\n  id, document_id, chunk_index, content, metadata_json, embedding, embedding_model, created_at')
h2('7.4  Quy tắc retrieval')
bullet('Chỉ truy xuất destination PUBLISHED.')
bullet('Áp dụng filter tỉnh, vùng, category, price và accessibility trước vector ranking.')
bullet('Dùng hybrid search thay vì chỉ vector similarity.')
bullet('Rerank top candidates rồi mới đưa vào prompt.')
bullet('Mọi citation phải trỏ về source ID và thời điểm cập nhật.')

h1('8  Embedding và tìm kiếm hybrid')
p('BGE M3 là ứng viên phù hợp vì hỗ trợ đa ngôn ngữ, dense retrieval, sparse retrieval và input dài. Tuy nhiên vẫn cần benchmark riêng với truy vấn tiếng Việt của VISTA.')
h2('8.1  Hybrid retrieval')
code('vector_score = cosine_similarity(query_embedding, chunk_embedding)\nkeyword_score = PostgreSQL full text rank\nfinal_score = reciprocal_rank_fusion(vector_rank, keyword_rank)\nresult = rerank(top_candidates)')
p('Vector search tìm nội dung gần nghĩa; lexical search giữ tốt tên riêng như Bà Nà, Hội An hoặc mã địa điểm. Kết hợp hai loại giảm bỏ sót và giảm kết quả gần nghĩa nhưng sai địa danh.')
h2('8.2  Index và cập nhật')
bullet('Dùng pgvector cùng PostgreSQL hiện tại để giảm số hệ thống phải vận hành.')
bullet('Chỉ thêm HNSW khi dữ liệu và tải đủ lớn; dữ liệu nhỏ có thể exact search tốt hơn.')
bullet('Tạo index concurrently ở production.')
bullet('Lưu embedding model và version để reindex có kiểm soát.')

h1('9  Tool calling và luật nghiệp vụ')
p('Tool calling biến LLM thành bộ điều phối có giới hạn. Model được phép yêu cầu dữ liệu hoặc phép tính qua tool, nhưng tool do hệ thống kiểm soát và validate.')
table(['Tool','Đầu vào','Đầu ra tin cậy'],[
('search_destinations','Tỉnh, category, budget, accessibility','Danh sách ID PUBLISHED có metadata.'),
('get_destination_details','Destination ID','Giá, giờ mở cửa, tọa độ, nguồn.'),
('calculate_route','Danh sách tọa độ đã kiểm tra','Khoảng cách, thời lượng và route.'),
('calculate_budget','Activity và thành viên','BigDecimal totals do backend tính.'),
('get_trip_constraints','Trip ID từ context bảo mật','Ngày, thành viên, giới hạn của chuyến đi.')],[1.65,2.35,2.35])
p('Không cung cấp tool chạy SQL, shell hoặc gọi URL tùy ý. Tool phải có allowlist, schema, timeout, authorization và audit.')

h1('10  Structured output và kiểm chứng')
h2('10.1  JSON Schema')
code('{\n  "type": "object",\n  "required": ["days", "estimatedTotalVnd"],\n  "properties": {\n    "days": {"type": "array", "minItems": 1, "maxItems": 7},\n    "estimatedTotalVnd": {"type": "integer", "minimum": 0}\n  },\n  "additionalProperties": false\n}')
h2('10.2  Hai tầng validation')
bullet('Tầng cú pháp kiểm tra JSON Schema, type, enum, độ dài và số phần tử.')
bullet('Tầng nghiệp vụ kiểm tra destination tồn tại, số ngày, giờ, tọa độ, ngân sách và quyền truy cập.')
p('Nếu output lỗi, chỉ cho phép một lần repair có giới hạn. Nếu vẫn lỗi, trả trạng thái FAILED rõ ràng. Không dựng dữ liệu mẫu rồi báo thành công.')

h1('11  Thiết kế AI service')
h2('11.1  Cấu trúc đề xuất')
code('ai-service/\n  app/\n    api/\n    inference/\n    retrieval/\n    prompts/\n    schemas/\n    tools/\n    guardrails/\n    evaluation/\n  tests/\n  Dockerfile\n  requirements.lock')
h2('11.2  API nội bộ')
code('POST /internal/v1/itineraries/generate\nAuthorization: Bearer <service-token>\nX-Request-ID: <uuid>\nIdempotency-Key: <uuid>')
p('API nội bộ không được public. Spring Boot truyền dữ liệu đã chuẩn hóa và không truyền JWT người dùng cho model service. Service token phải được xoay vòng và lưu bằng secret manager.')
h2('11.3  Response có khả năng truy vết')
code('{\n  "model": "qwen3-8b",\n  "modelRevision": "pinned-revision",\n  "promptVersion": "itinerary-v1",\n  "retrievalVersion": "destinations-v1",\n  "days": [],\n  "citations": [],\n  "warnings": []\n}')

h1('12  Tích hợp Spring Boot')
h2('12.1  Provider abstraction')
code('public interface AiProvider {\n  AiPlanResult generate(AiPlanCommand command);\n  boolean isHealthy();\n  String providerName();\n}')
p('Tạo LocalLlmProvider và giữ GeminiProvider trong thời gian chuyển đổi. AiProviderRouter chọn provider bằng feature flag, health, quota và rollout percentage. Không đặt fallback âm thầm; response phải ghi rõ provider và fallback status cho telemetry.')
h2('12.2  Luồng service')
for x in ['CurrentUserProvider lấy danh tính.','TripAccessService xác minh thành viên.','Request DTO được validate và chuẩn hóa.','Quota service thực hiện atomic consume.','Tạo AI job và idempotency record.','Worker gọi AI Orchestrator.','Kết quả được validate lại trong Java.','Lưu preview; chỉ commit itinerary sau xác nhận của người dùng.']:
    number(x)
h2('12.3  Error contract')
table(['Mã','HTTP','Ý nghĩa'],[
('AI_RATE_LIMITED','429','Hết quota, có retryAfter.'),('AI_BUSY','503','Hàng đợi đầy.'),('AI_TIMEOUT','504','Model quá thời gian.'),('AI_INVALID_OUTPUT','502','Model trả dữ liệu không hợp lệ.'),('AI_RETRIEVAL_EMPTY','422','Không đủ dữ liệu đáng tin cậy.'),('AI_JOB_CANCELLED','409','Job đã bị hủy.')],[1.7,0.8,4.0])

h1('13  Hàng đợi và streaming')
p('Sinh lịch trình có thể kéo dài hơn request web thông thường. Nên dùng job bất đồng bộ thay vì giữ connection đến khi model kết thúc.')
code('POST /trips/{id}/ai/jobs  -> 202 + jobId\nGET  /ai/jobs/{jobId}        -> QUEUED | RUNNING | COMPLETED | FAILED\nDELETE /ai/jobs/{jobId}      -> yêu cầu hủy\nGET  /ai/jobs/{jobId}/events -> SSE cho tiến trình')
bullet('Queue giới hạn số job chạy đồng thời theo GPU.')
bullet('Admission control trả AI_BUSY khi queue vượt ngưỡng.')
bullet('Idempotency ngăn double click tạo nhiều job.')
bullet('Timeout và cancellation phải giải phóng slot inference.')

h1('14  Bảo mật và quyền riêng tư')
h2('14.1  Mô hình đe dọa')
table(['Rủi ro','Ví dụ','Biện pháp'],[
('Prompt injection','Tên chuyến đi yêu cầu bỏ qua luật.','Tách system và data, tool allowlist, validate output.'),
('Cross trip leakage','RAG lấy dữ liệu chuyến đi khác.','Filter bằng tenant và trip scope trước retrieval.'),
('SSRF','Model gọi URL do người dùng cung cấp.','Không có generic HTTP tool.'),
('Denial of wallet','Spam inference làm quá tải GPU.','Rate limit, quota, queue và max tokens.'),
('Sensitive logging','Prompt chứa email hoặc tọa độ.','Redaction và log metadata thay vì nội dung.'),
('Model endpoint exposure','Internet gọi thẳng vLLM.','Private network và reverse proxy.')],[1.45,2.4,2.65])
h2('14.2  Dữ liệu gửi vào model')
bullet('Chỉ gửi dữ liệu tối thiểu cần cho nhiệm vụ.')
bullet('Không gửi access token, refresh token, mật khẩu hoặc service secret.')
bullet('Không gửi vị trí thời gian thực nếu lập lịch trình không cần.')
bullet('Gắn retention policy cho prompt và response.')
bullet('Cho phép người dùng xóa lịch sử AI liên quan tài khoản.')
h2('14.3  Network')
bullet('Model server không có public ingress.')
bullet('Chặn outbound Internet từ model container nếu không cần.')
bullet('Chỉ Orchestrator được phép gọi model server.')
bullet('Bảo vệ cả endpoint không nằm dưới /v1 bằng reverse proxy.')

h1('15  Quan sát và kiểm soát chi phí')
h2('15.1  Metrics')
table(['Nhóm','Metric cần có'],[
('Traffic','requests, jobs, concurrency, queue depth.'),('Latency','TTFT, total latency, retrieval latency, validation latency.'),('Tokens','input, output, cached tokens, tokens per second.'),('Quality','schema failures, repair rate, fallback rate, user acceptance.'),('Resources','GPU utilization, VRAM, power, temperature, OOM.'),('Reliability','timeouts, cancellations, provider errors, circuit state.')],[1.5,5.0])
h2('15.2  Log và tracing')
p('Mỗi request cần request ID, job ID, user hash, trip ID hash, model revision, prompt version, retrieval version và timing. Không log raw prompt mặc định. Trace phải nối từ client qua Spring Boot, Orchestrator, retrieval và model server.')
h2('15.3  Tính giá nội bộ')
p('Self host không miễn phí. Chi phí gồm GPU, điện, storage, bandwidth, nhân sự vận hành và công suất nhàn rỗi. Theo dõi chi phí trên mỗi job thành công, không chỉ chi phí trên mỗi token.')

h1('16  Bộ đánh giá chất lượng')
h2('16.1  Dataset')
p('Tạo 200 đến 500 case tiếng Việt, có kết quả mong đợi hoặc rubric. Chia thành development set và locked test set để tránh tối ưu quá mức.')
bullet('Chuyến đi 1 đến 7 ngày và nhiều mức ngân sách.')
bullet('Gia đình, trẻ nhỏ, người cao tuổi và nhu cầu accessibility.')
bullet('Mùa mưa, địa điểm đóng cửa, dữ liệu thiếu hoặc mâu thuẫn.')
bullet('Prompt injection, nội dung độc hại và yêu cầu truy cập chuyến đi khác.')
bullet('Tọa độ sai, giá âm, destination không tồn tại và output quá dài.')
h2('16.2  Chỉ số')
table(['Chỉ số','Ngưỡng khởi điểm đề xuất'],[
('JSON schema pass rate','Tối thiểu 99 phần trăm sau tối đa một repair.'),('Grounded destination rate','100 phần trăm destination ID tồn tại và PUBLISHED.'),('Budget violation','0 phần trăm sau backend validation.'),('Cross trip leakage','0 case.'),('P95 latency','Đặt theo UX thực tế sau benchmark.'),('User acceptance','Theo dõi tỷ lệ lưu hoặc chỉnh sửa kế hoạch.')],[2.2,4.3])
h2('16.3  Shadow traffic')
p('Trong giai đoạn chuyển đổi, gửi cùng request đã ẩn danh sang Gemini và local model nhưng chỉ trả một kết quả cho người dùng. So sánh tự động và review mẫu. Không lưu hai bộ dữ liệu nhạy cảm lâu hơn cần thiết.')

h1('17  Fine tuning và LoRA')
h2('17.1  Khi nào cần')
p('Chỉ fine tune khi RAG và prompt đã ổn nhưng model vẫn sai về phong cách, format hoặc quy trình lặp lại. Không fine tune để ghi nhớ giá vé hoặc giờ mở cửa.')
h2('17.2  Dữ liệu huấn luyện')
bullet('Chỉ dùng dữ liệu có quyền sử dụng.')
bullet('Loại bỏ PII và dữ liệu chuyến đi riêng tư.')
bullet('Dùng ví dụ đã được chuyên gia duyệt.')
bullet('Giữ test set tách biệt tuyệt đối.')
bullet('Theo dõi dataset version, base model và adapter checksum.')
h2('17.3  Quy trình LoRA')
for x in ['Xác định lỗi hành vi lặp lại bằng evaluation.','Chuẩn bị instruction response pairs chất lượng cao.','Train adapter LoRA trên base model đã pin.','So sánh với base model trên locked test set.','Kiểm tra regression, safety và tiếng Việt.','Canary rollout; có khả năng rollback adapter.']:
    number(x)

h1('18  Quy trình triển khai từng giai đoạn')
h2('Giai đoạn 0  Chuẩn bị')
bullet('Định nghĩa SLO, budget, privacy policy và tiêu chí thay Gemini.')
bullet('Tạo evaluation dataset ban đầu.')
bullet('Thêm AiProvider abstraction và telemetry.')
h2('Giai đoạn 1  Proof of concept')
bullet('Dựng llama.cpp và Qwen3 8B GGUF trên máy phát triển.')
bullet('Kết nối qua OpenAI compatible endpoint.')
bullet('Chạy 100 prompt, đo chất lượng và latency.')
h2('Giai đoạn 2  RAG')
bullet('Cài pgvector, tạo schema knowledge và ingestion job.')
bullet('Dựng BGE M3 embedding service.')
bullet('Triển khai hybrid retrieval và citations.')
h2('Giai đoạn 3  Production serving')
bullet('Dựng vLLM trên GPU server private.')
bullet('Thêm reverse proxy, queue, quota, circuit breaker và metrics.')
bullet('Load test, chaos test và kiểm tra OOM recovery.')
h2('Giai đoạn 4  Chuyển traffic')
bullet('Shadow traffic rồi canary 5 phần trăm.')
bullet('Tăng 10, 25, 50 và 100 phần trăm khi metric đạt ngưỡng.')
bullet('Giữ Gemini fallback có thời hạn và alert mỗi lần sử dụng.')
h2('Giai đoạn 5  Tối ưu')
bullet('Prompt caching, quantization và batching.')
bullet('Chỉ cân nhắc LoRA sau khi có dữ liệu thực tế.')

h1('19  Runbook vận hành')
h2('19.1  Model server OOM')
number('Ngừng nhận job mới và mở circuit breaker.')
number('Giữ job đang chờ, trả trạng thái AI_BUSY.')
number('Thu thập VRAM, batch, context và request gây lỗi.')
number('Restart có kiểm soát; không retry vô hạn.')
number('Giảm concurrency hoặc context trước khi mở traffic.')
h2('19.2  Chất lượng giảm sau nâng model')
number('Dừng rollout và quay về model revision trước.')
number('So sánh prompt version, tokenizer, chat template và quantization.')
number('Chạy locked regression suite.')
number('Chỉ rollout lại sau khi có báo cáo nguyên nhân.')
h2('19.3  RAG trả dữ liệu sai')
number('Tắt index hoặc source revision lỗi.')
number('Kiểm tra filter PUBLISHED, tenant scope và content hash.')
number('Reindex có version; không overwrite mù.')
number('Kiểm tra citation trước khi mở lại.')
h2('19.4  Key hoặc model bị xâm phạm')
number('Cô lập model service và thu hồi service token.')
number('Chặn network, bảo toàn log metadata và xác định phạm vi.')
number('Thay weights từ checksum tin cậy và xoay secret.')
number('Thông báo theo incident policy nếu có dữ liệu người dùng.')

h1('20  Checklist nghiệm thu')
h2('Kiến trúc')
for x in ['AI Provider abstraction hoạt động.','Model server private và không public ingress.','RAG chỉ dùng dữ liệu PUBLISHED.','Backend là nơi validation và lưu nghiệp vụ.']:
    bullet('☐ '+x)
h2('Bảo mật')
for x in ['Không log prompt chứa PII.','Không có generic SQL hoặc HTTP tool.','Rate limit và quota atomic.','Có audit và request tracing.','Có test cross trip leakage.']:
    bullet('☐ '+x)
h2('Chất lượng')
for x in ['Schema pass đạt ngưỡng.','Không destination bịa.','Không vượt ngân sách sau validation.','Có citations.','Locked evaluation suite chạy trong CI.']:
    bullet('☐ '+x)
h2('Vận hành')
for x in ['Có dashboard GPU và queue.','Có timeout, cancellation và circuit breaker.','Có canary và rollback.','Weights, prompt, embedding đều có version.','Runbook đã diễn tập.']:
    bullet('☐ '+x)

h1('Phụ lục A  Cấu hình khởi đầu đề xuất')
code('Model: Qwen3 8B\nDevelopment: llama.cpp với GGUF Q4 hoặc Q5\nProduction: vLLM trên GPU server private\nEmbedding: BGE M3\nVector store: PostgreSQL và pgvector\nQueue và quota: Redis\nOrchestration: service riêng\nBusiness validation: Spring Boot\nClient progress: job polling hoặc SSE')
h1('Phụ lục B  Câu hỏi cần trả lời trước khi mua hạ tầng')
for x in ['Bao nhiêu request AI mỗi ngày và vào giờ cao điểm?','Context và output trung bình bao nhiêu token?','P95 latency chấp nhận được là bao nhiêu?','Có bao nhiêu request đồng thời?','Dữ liệu nào được phép đi vào model?','SLO và thời gian khôi phục là bao nhiêu?','Ngân sách GPU theo tháng là bao nhiêu?','Ngưỡng chất lượng nào cho phép tắt Gemini?']:
    number(x)
h1('Phụ lục C  Nguồn tham khảo chính')
refs=[
('Qwen3 official announcement','https://qwenlm.github.io/blog/qwen3/'),
('Qwen3 8B model card','https://huggingface.co/Qwen/Qwen3-8B'),
('vLLM documentation','https://docs.vllm.ai/en/stable/'),
('vLLM OpenAI compatible server and security note','https://github.com/vllm-project/vllm/blob/main/docs/serving/online_serving/openai_compatible_server.md'),
('llama.cpp official repository','https://github.com/ggml-org/llama.cpp'),
('BGE M3 model card','https://huggingface.co/BAAI/bge-m3'),
('pgvector official repository','https://github.com/pgvector/pgvector')]
for name,url in refs: p(f'{name}: {url}')

# Footer with page field
for section in doc.sections:
    fp=section.footer.paragraphs[0]
    fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=fp.add_run('VISTA  Cẩm nang LLM tự host   |   ')
    run.font.size=Pt(8); run.font.color.rgb=RGBColor(90,90,90)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld)

doc.core_properties.title='Cẩm nang triển khai LLM tự host cho hệ thống VISTA'
doc.core_properties.subject='Kiến thức và hướng dẫn triển khai nền tảng AI tự host'
doc.core_properties.author='VISTA Engineering'
doc.save(OUT)
print(OUT)
